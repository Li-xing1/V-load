import os
import re
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.utils.dataframe import dataframe_to_rows
from scipy.stats import ks_2samp, wasserstein_distance


def _finite(values):
    values = np.asarray(values, dtype=np.float64).reshape(-1)
    return values[np.isfinite(values)]


def reference_bin_edges(reference, bins=50):
    reference = _finite(reference)
    quantiles = np.linspace(0.0, 1.0, min(int(bins), reference.size) + 1)
    edges = np.unique(np.quantile(reference, quantiles))
    if edges.size < 2:
        value = edges[0]
        edges = np.array([value - 0.5, value + 0.5])
    edges[0], edges[-1] = -np.inf, np.inf
    return edges


def normalized_js_divergence(real, generated, bin_edges=None):
    real, generated = _finite(real), _finite(generated)
    if real.size == 0 or generated.size == 0:
        raise ValueError("Distributions must be non-empty after removing non-finite values.")
    edges = reference_bin_edges(real) if bin_edges is None else np.asarray(bin_edges)
    p = np.histogram(real, bins=edges)[0].astype(np.float64)
    q = np.histogram(generated, bins=edges)[0].astype(np.float64)
    p, q = p / p.sum(), q / q.sum()
    midpoint = 0.5 * (p + q)
    p_term = np.zeros_like(p)
    q_term = np.zeros_like(q)
    p_valid, q_valid = p > 0, q > 0
    p_term[p_valid] = p[p_valid] * np.log2(p[p_valid] / midpoint[p_valid])
    q_term[q_valid] = q[q_valid] * np.log2(q[q_valid] / midpoint[q_valid])
    return float(0.5 * (p_term.sum() + q_term.sum()))


def continuous_metrics(real, generated, bin_edges=None):
    real, generated = _finite(real), _finite(generated)
    if real.size == 0 or generated.size == 0:
        raise ValueError("Distributions must be non-empty after removing non-finite values.")
    return {
        "wasserstein_1": float(wasserstein_distance(real, generated)),
        "ks_statistic": float(ks_2samp(real, generated).statistic),
        "js_divergence": normalized_js_divergence(real, generated, bin_edges),
    }


def quantile_errors(real, generated, quantiles=(0.5, 0.95, 0.99)):
    real, generated = _finite(real), _finite(generated)
    return {
        f"p{int(q * 100):02d}_error": float(abs(np.quantile(generated, q) - np.quantile(real, q)))
        for q in quantiles
    }


def range_violation_rate(values, low=None, high=None):
    values = _finite(values)
    if values.size == 0:
        raise ValueError("Values must contain at least one finite observation.")
    invalid = np.zeros(values.shape, dtype=bool)
    if low is not None:
        invalid |= values < low
    if high is not None:
        invalid |= values > high
    return float(invalid.mean())


def crps_ensemble(samples, real):
    samples = np.asarray(samples, dtype=np.float64)
    real = np.asarray(real, dtype=np.float64)
    if samples.shape[0] != real.shape[0]:
        raise ValueError("CRPS samples and observations must have the same condition axis.")
    term_observation = np.mean(np.abs(samples - np.expand_dims(real, axis=1)))
    pairwise = np.abs(np.expand_dims(samples, 2) - np.expand_dims(samples, 1))
    return float(term_observation - 0.5 * pairwise.mean())


def energy_score(samples, real):
    samples = np.asarray(samples, dtype=np.float64)
    real = np.asarray(real, dtype=np.float64)
    samples = samples.reshape(samples.shape[0], samples.shape[1], -1)
    real = real.reshape(real.shape[0], -1)
    first = np.linalg.norm(samples - real[:, None, :], axis=-1).mean()
    pairwise = np.linalg.norm(samples[:, :, None, :] - samples[:, None, :, :], axis=-1).mean()
    return float(first - 0.5 * pairwise)


def categorical_metrics(real_counts, generated_counts, rare_threshold=0.05):
    real = np.asarray(real_counts, dtype=np.float64)
    generated = np.asarray(generated_counts, dtype=np.float64)
    real, generated = real / real.sum(), generated / generated.sum()
    midpoint = 0.5 * (real + generated)
    js = 0.0
    for dist in (real, generated):
        valid = dist > 0
        js += 0.5 * np.sum(dist[valid] * np.log2(dist[valid] / midpoint[valid]))
    errors = np.abs(real - generated)
    rare = real < rare_threshold
    return {
        "js_divergence": float(js),
        "total_variation": float(0.5 * errors.sum()),
        "class_proportion_mae": float(errors.mean()),
        "rare_class_mae": float(errors[rare].mean()) if rare.any() else np.nan,
    }


def threshold_probability_errors(real, generated, thresholds=(5.0, 20.0)):
    real, generated = _finite(real), _finite(generated)
    return {
        f"probability_below_{threshold:g}_error": float(abs(np.mean(generated < threshold) - np.mean(real < threshold)))
        for threshold in thresholds
    }


def correlation_mae(real, generated):
    real_corr = np.corrcoef(np.asarray(real), rowvar=False)
    generated_corr = np.corrcoef(np.asarray(generated), rowvar=False)
    return float(np.nanmean(np.abs(real_corr - generated_corr)))


def _sheet_name(model, variant):
    name = re.sub(r"[^A-Za-z0-9_]+", "_", f"{model}_{variant or 'overall'}")
    return name[:31]


def reset_model_evaluation(path, model):
    path = Path(path)
    if not path.exists():
        return
    summary = pd.read_excel(path, sheet_name="summary")
    summary = summary.loc[summary["model"] != model]
    workbook = load_workbook(path)
    for sheet_name in workbook.sheetnames[:]:
        if sheet_name == "summary" or sheet_name.startswith(f"{model}_"):
            del workbook[sheet_name]
    summary_sheet = workbook.create_sheet("summary", 0)
    for row in dataframe_to_rows(summary, index=False, header=True):
        summary_sheet.append(row)
    workbook.save(path)


def update_evaluation_report(path, model, variant, metrics, units=None):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    units = units or {}
    rows = pd.DataFrame(
        [{"model": model, "variant": variant or "overall", "metric": key,
          "value": float(value), "unit": units.get(key, "")} for key, value in metrics.items()]
    )
    if path.exists():
        summary = pd.read_excel(path, sheet_name="summary")
        keep = ~((summary["model"] == model) & (summary["variant"] == (variant or "overall")))
        summary = pd.concat([summary.loc[keep], rows], ignore_index=True)
        workbook = load_workbook(path)
    else:
        summary = rows.copy()
        workbook = Workbook()
        del workbook[workbook.active.title]

    variant_sheet_name = _sheet_name(model, variant)
    for sheet_name in ("summary", variant_sheet_name):
        if sheet_name in workbook.sheetnames:
            del workbook[sheet_name]
    summary_sheet = workbook.create_sheet("summary", 0)
    variant_sheet = workbook.create_sheet(variant_sheet_name)
    for row in dataframe_to_rows(summary, index=False, header=True):
        summary_sheet.append(row)
    for row in dataframe_to_rows(rows, index=False, header=True):
        variant_sheet.append(row)
    workbook.save(path)
    return os.fspath(path)

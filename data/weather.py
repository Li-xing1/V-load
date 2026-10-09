import json, urllib
from urllib.parse import urlencode
import requests
import openpyxl
import os
from tqdm import tqdm
import pandas as pd
import pickle
import argparse
from datetime import datetime, timedelta


BASE_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_DIR = os.path.join(BASE_DIR, 'jsq')
WEATHER_PKL = os.path.join(OUTPUT_DIR, 'HEBWeather.pkl')
WEATHER_XLSX = os.path.join(OUTPUT_DIR, 'Weather.xlsx')
CITY_XLSX = os.path.join(OUTPUT_DIR, 'city.xlsx')


def ensure_output_dir():
    os.makedirs(OUTPUT_DIR, exist_ok=True)


def iter_dates(start_date, end_date):
    start = datetime.strptime(start_date, '%Y%m%d')
    end = datetime.strptime(end_date, '%Y%m%d')
    if start > end:
        raise ValueError('start_date must be earlier than or equal to end_date')

    current = start
    while current <= end:
        yield current.strftime('%Y%m%d')
        current += timedelta(days=1)


def append_df_to_excel(filename, df):
    key = list(df.keys())
    row = ['A', 'B', 'C', 'D', 'E', 'F', 'G', 'H', 'I', 'J', 'K', 'L', 'M', 'N', 'O', 'P', 'Q', 'R', 'S', 'T', 'U', 'V',
           'W', 'X', 'Y', 'Z', 'AA', 'AB', 'AC', 'AD', 'AE', 'AF', 'AG', 'AH', 'AI', 'AJ', 'AK', 'AL', 'AM', 'AN']
    if not os.path.exists(filename):
        wb = openpyxl.Workbook()
        sheet = wb.active
        for i in range(len(key)):
            sheet[row[i] + '1'] = key[i]
    else:
        wb = openpyxl.load_workbook(filename)
        sheet = wb.active
    row = sheet.max_row + 1
    for i in range(len(key)):
        sheet.cell(row=row, column=i + 1, value=df[key[i]])
    wb.save(filename)
    wb.close()


def export_to_excel(data, file_path):
    # 检查Excel文件是否存在
    if os.path.exists(file_path):
        # 如果文件存在，直接使用pandas的DataFrame读取已有的Excel文件
        df = pd.read_excel(file_path)
    else:
        # 如果文件不存在，创建一个空的DataFrame
        df = pd.DataFrame()
    # 将字典列表转换为DataFrame
    new_data = pd.DataFrame(data)
    # 合并新的数据到旧的数据中
    df = pd.concat([df, new_data], ignore_index=True)
    # 将数据保存到Excel文件中
    df.to_excel(file_path, index=False)


def history_weather(APPKEY, SIGN, weaId, start_date='20240601', end_date='20240705'):
    url = 'http://api.k780.com'
    date_list = list(iter_dates(start_date, end_date))

    date_list1 = []
    for date in date_list:
        print(date)
        params = {
            'app': 'weather.history',
            'weaId': weaId,
            'date': date,
            'appkey': APPKEY,
            'sign': SIGN,
            'format': 'json',
        }
        response = requests.get(url, params=params)
        a_result = response.json()

        if a_result:
            if a_result['success'] != '0':
                dtlist = a_result['result']
                print(dtlist)
                date_list1 += dtlist

            else:
                print(a_result['msgid'] + ' ' + a_result['msg'])
        else:
            print('Request nowapi fail.')
    ensure_output_dir()
    with open(WEATHER_PKL, 'wb') as f:
        pickle.dump(date_list1, f)


def cityid(APPKEY, SIGN):
    url = 'http://api.k780.com'
    params = {
        'app': 'weather.city',
        'areaType': 'cn',
        'appkey': APPKEY,
        'sign': SIGN,
        'format': 'json',
    }
    response = requests.get(url, params=params)
    a_result = response.json()

    if a_result:
        if a_result['success'] != '0':
            dtlist = a_result['result']['dtList']
            keys = dtlist.keys()
            ensure_output_dir()
            for key in tqdm(keys):
                append_df_to_excel(CITY_XLSX, dtlist[key])

        else:
            print(a_result['msgid'] + ' ' + a_result['msg'])
    else:
        print('Request nowapi fail.')


def pkl2excel():
    with open(WEATHER_PKL, 'rb') as f:
        data = pickle.load(f)
    df = pd.DataFrame(columns=['uptime', 'weatid', 'temp', 'humidity', 'winpid', 'aqi'])
    for datai in data:
        df.loc[len(df)] = [datai[i] for i in df.columns]
        df.loc[len(df) - 1, 'humidity'] = df.loc[len(df) - 1, 'humidity'].replace('%', '')
    df['uptime'] = pd.to_datetime(df['uptime'])
    df['weatid'] = df['weatid'].astype(int)
    df['temp'] = df['temp'].astype(int)
    df['humidity'] = df['humidity'].astype(int)
    df['winpid'] = df['winpid'].astype(int)
    df['aqi'] = df['aqi'].astype(int)
    ensure_output_dir()
    df.to_excel(WEATHER_XLSX, index=False)


def parse_args():
    parser = argparse.ArgumentParser()
    parser.add_argument('--start-date', default='20260710', help='Start date, format YYYYMMDD')
    parser.add_argument('--end-date', default='20260917', help='End date, format YYYYMMDD')
    parser.add_argument('--wea-id', default='169', help='Weather city id')
    return parser.parse_args()


if __name__ == '__main__':
    args = parse_args()
    APPKEY = os.environ.get('K780_APPKEY')
    SIGN = os.environ.get('K780_SIGN')
    if not APPKEY or not SIGN:
        raise SystemExit('Please set K780_APPKEY and K780_SIGN environment variables before requesting weather data.')
    history_weather(APPKEY, SIGN, weaId=args.wea_id, start_date=args.start_date, end_date=args.end_date)
    # cityid(APPKEY, SIGN)
    pkl2excel()
